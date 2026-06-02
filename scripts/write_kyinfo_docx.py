# -*- coding: utf-8 -*-
"""Generate KyInfo interface documentation into D:\\KyInfo.docx"""

from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

SCRIPT_DIR = Path(__file__).resolve().parent
SCREENSHOTS_DIR = SCRIPT_DIR.parent / "docs" / "screenshots"


def set_cell_text(cell, text):
    cell.text = text
    for p in cell.paragraphs:
        for r in p.runs:
            r.font.size = Pt(10.5)
            r.font.name = "宋体"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            set_cell_text(table.rows[ri + 1].cells[ci], str(val))
    doc.add_paragraph()


def add_screenshots(doc, filenames, caption="界面截图", note=None):
    doc.add_heading(caption, level=2)
    if note:
        doc.add_paragraph(note)
    if not filenames:
        doc.add_paragraph("（暂无对应截图，可在系统运行后访问上述路由自行截取补充。）")
        return
    for idx, name in enumerate(filenames):
        path = SCREENSHOTS_DIR / name
        if not path.is_file():
            doc.add_paragraph(f"（未找到截图文件：{name}）")
            continue
        if len(filenames) > 1:
            cap = doc.add_paragraph(f"{caption}（{idx + 1}/{len(filenames)}）")
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_picture(str(path), width=Cm(15.5))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()


def add_section(
    doc,
    fig_no,
    title,
    route,
    flow_lines,
    mock_title,
    mock_table_headers,
    mock_table_rows,
    db_sections,
    screenshots=None,
    screenshot_note=None,
):
    doc.add_heading(f"{fig_no} {title}", level=1)
    if screenshots is not None:
        add_screenshots(doc, screenshots, note=screenshot_note)
    p = doc.add_paragraph()
    r = p.add_run(f"路由：{route}")
    r.bold = True
    doc.add_heading("功能与流程", level=2)
    for line in flow_lines:
        doc.add_paragraph(line, style="List Number")
    if mock_table_headers:
        doc.add_heading(mock_title or "模拟数据", level=2)
        add_table(doc, mock_table_headers, mock_table_rows)
    if db_sections:
        doc.add_heading("数据库表", level=2)
        for db_name, fields in db_sections:
            doc.add_paragraph(db_name, style="List Bullet")
            if isinstance(fields, list) and fields and isinstance(fields[0], tuple):
                add_table(doc, ["字段名", "说明"], fields)
            elif isinstance(fields, list) and fields and isinstance(fields[0], str):
                add_table(doc, ["字段名", "说明"], [(f, "") for f in fields])
            else:
                doc.add_paragraph(str(fields))


def main():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(12)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    title = doc.add_heading("KyInfo 考研信息系统 — 界面说明文档", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(
        "本文档描述系统各功能界面的编号、名称、业务流程、模拟数据及关联数据库表字段。"
        "表名与项目 EF Core 映射一致（如 Users、Schools 等）。"
    )

    user_fields = [
        ("Id", "主键"),
        ("UserName", "用户名（登录标识）"),
        ("Email", "邮箱（登录标识）"),
        ("PasswordHash", "密码哈希"),
        ("Role", "角色：User(0)、Admin(1)、Root(2)"),
        ("CreatedAt", "创建时间（UTC）"),
    ]
    exam_fields = [
        ("Id", "主键"),
        ("Year", "考试年份"),
        ("TotalScore", "总分（必填）"),
        ("PoliticsScore", "政治（可选）"),
        ("EnglishScore", "英语（可选）"),
        ("MathScore", "数学（可选）"),
        ("MajorSubjectScore", "专业课（可选）"),
        ("UserId", "外键 → Users.Id"),
        ("SchoolId", "外键 → Schools.Id（可选）"),
        ("MajorId", "外键 → Majors.Id（可选）"),
        ("CreatedAt", "创建时间"),
    ]
    school_fields = [
        ("Id", "主键"),
        ("Name", "院校全称"),
        ("ShortName", "简称"),
        ("Province", "省份"),
        ("City", "城市"),
        ("LevelTag", "层次（985/211/普通等）"),
        ("Type", "类型（综合、理工等）"),
        ("Property", "办学性质（公办等）"),
        ("Website", "官网"),
        ("CreatedAt", "创建时间"),
    ]
    major_fields = [
        ("Id", "主键"),
        ("Name", "专业名称"),
        ("Code", "专业代码"),
        ("DisciplineCategory", "学科门类"),
        ("DegreeType", "学位类型（学硕/专硕）"),
        ("StudyType", "学习方式"),
        ("DurationYears", "学制（年）"),
        ("TuitionPerYear", "年学费"),
        ("SchoolDepartment", "院系"),
        ("Description", "说明"),
        ("SchoolId", "外键 → Schools.Id"),
    ]
    scoreline_fields = [
        ("Id", "主键"),
        ("Year", "年份"),
        ("Score", "分数线"),
        ("IsNational", "是否国家线"),
        ("SchoolId", "院校 Id（国家线为空）"),
        ("MajorId", "专业 Id（国家线为空）"),
        ("Note", "备注"),
        ("CreatedAt", "创建时间"),
    ]
    recruit_fields = [
        ("Id", "主键"),
        ("Year", "招生年份"),
        ("SchoolId", "院校 Id"),
        ("MajorId", "专业 Id"),
        ("PlanCount", "计划招生人数"),
        ("ExamSubjects", "考试科目"),
        ("ExtraRequirements", "其他要求"),
        ("SourceUrl", "原文链接"),
        ("PublishedAt", "发布时间"),
        ("CreatedAt", "创建时间"),
    ]
    audit_fields = [
        ("Id", "主键"),
        ("CreatedAtUtc", "操作时间"),
        ("ActorUserId", "操作人 Id"),
        ("ActorRole", "操作人角色"),
        ("Action", "动作（如 User.Create、ExamScore.ImportBatch）"),
        ("ResourceType", "资源类型"),
        ("ResourceId", "资源 Id"),
        ("Summary", "摘要"),
    ]

    # 图1-1
    add_section(
        doc,
        "图1-1",
        "系统登录界面图",
        "/login",
        [
            "未登录用户进入登录页，输入用户名或邮箱及密码。",
            "点击「登录」，前端调用认证 API，后端校验 Users 表中的 PasswordHash。",
            "成功则写入会话（JWT/浏览器会话）并跳转首页或原目标页；失败则提示错误。",
            "点击「去注册」跳转图1-2。",
        ],
        "模拟数据",
        ["字段", "模拟值"],
        [
            ("用户名或邮箱", "zhangsan@test.com 或 张三"),
            ("密码", "Test@123456（以实际部署为准）"),
        ],
        [("Users", user_fields)],
        screenshots=["fig1-1-login.png"],
    )

    # 图1-2
    add_section(
        doc,
        "图1-2",
        "系统注册界面图",
        "/register",
        [
            "填写用户名、邮箱、密码。",
            "点击「注册」，后端检查用户名/邮箱唯一性。",
            "密码哈希后插入 Users，默认 Role = User。",
            "成功后跳转登录或自动登录。",
        ],
        "模拟数据",
        ["字段", "模拟值"],
        [
            ("用户名", "StudyMaster_2026"),
            ("邮箱", "student_test@example.com"),
            ("密码", "SecurePass123!"),
        ],
        [("Users", user_fields)],
        screenshots=["fig1-2-register.png"],
    )

    # 图1-3
    add_section(
        doc,
        "图1-3",
        "系统首页界面图",
        "/",
        [
            "登录后展示欢迎信息与快捷入口（院校查询、分数线、智能志愿推荐等）。",
            "点击按钮跳转对应功能模块。",
            "本页不直接读写业务数据表。",
        ],
        None,
        None,
        None,
        [("无", "本界面不涉及数据库表")],
        screenshots=[],
    )

    # 图1-4
    add_section(
        doc,
        "图1-4",
        "我的档案界面图",
        "/account",
        [
            "基本资料：展示并修改当前用户的用户名、邮箱，点击「保存资料」更新 Users。",
            "修改密码：校验当前密码后更新 PasswordHash。",
            "我的成绩：按当前登录用户 Id 查询 ExamScores；可按年份筛选；填写后点击「新增成绩」。",
        ],
        "模拟数据",
        ["区域/项", "模拟值"],
        [
            ("用户名", "111"),
            ("邮箱", "2004@qq.com"),
            ("角色", "User"),
            ("注册时间", "2026-05-12 14:15（UTC）"),
            ("新增成绩", "年份 2025，总分 360，院校 Id 1，专业 Id 1"),
        ],
        [("Users", user_fields), ("ExamScores", exam_fields)],
        screenshots=["fig1-4-account.png"],
    )

    # 图1-5
    add_section(
        doc,
        "图1-5",
        "院校查询界面图",
        "/schools",
        [
            "输入关键词（名称/简称）、省份、层次标签（如 985、211）。",
            "点击「开始查询」，从 Schools 检索列表。",
            "点击列表项，右侧展示院校详情（简介、官网等）。",
        ],
        "模拟数据（查询结果示例）",
        ["名称", "简称", "地区", "标签"],
        [
            ("北京大学", "北大", "北京 / 北京", "985，综合·公办"),
            ("清华大学", "清华", "北京 / 北京", "985，综合·公办"),
            ("深圳大学", "深大", "广东 / 深圳", "普通，综合·公办"),
            ("南京师范大学", "南师大", "江苏 / 南京", "211，师范·公办"),
        ],
        [("Schools", school_fields)],
        screenshots=["fig1-5-schools.png"],
    )

    # 图1-6
    add_section(
        doc,
        "图1-6",
        "专业查询界面图",
        "/majors",
        [
            "输入关键词（名称/代码）及可选院校 Id。",
            "点击「开始查询」，关联 Majors 与 Schools。",
            "点击列表项，右侧展示学制、学费、说明等详情。",
        ],
        "模拟数据",
        ["专业", "代码", "院校", "标签"],
        [
            ("计算机科学与技术", "081200", "北京大学", "学硕，全日制，工学"),
            ("软件工程", "083500", "北京大学", "学硕，全日制，工学"),
            ("学科教学（语文）", "045103", "南京师范大学", "专硕，全日制，教育学"),
        ],
        [
            ("Majors", major_fields),
            ("Schools", [("Id", "主键"), ("Name", "院校名称（关联展示）")]),
        ],
        screenshots=["fig1-6-majors.png"],
    )
    doc.add_paragraph("筛选示例：关键词「计算机」，院校 Id「1」（北京大学）")

    # 图1-7
    add_section(
        doc,
        "图1-7",
        "分数线查询界面图",
        "/scorelines",
        [
            "可选填院校 Id、专业 Id、年份，选择类型（全部/国家线/院校专业线）。",
            "点击「查询」，从 ScoreLines 读取并表格展示。",
            "支持与国家线、各校专业复试线对比分析。",
        ],
        "模拟数据",
        ["年份", "分数", "类型", "院校", "专业"],
        [
            ("2025", "270", "国家线", "—", "—"),
            ("2025", "358", "院校/专业线", "北京大学", "计算机科学与技术"),
            ("2025", "338", "院校/专业线", "华中科技大学", "计算机科学与技术"),
            ("2025", "362", "院校/专业线", "南京师范大学", "教育学"),
        ],
        [("ScoreLines", scoreline_fields)],
        screenshots=["fig1-7-scorelines.png"],
    )

    # 图1-8
    add_section(
        doc,
        "图1-8",
        "分数线趋势界面图",
        "/scoreline-trends",
        [
            "设置院校 Id、专业 Id、类型（国家线/院校线）。",
            "点击「绘制趋势图」，按年份聚合 ScoreLines 数据。",
            "以折线图展示近年走势，辅助志愿填报判断。",
        ],
        "模拟数据（北大计科学硕 SchoolId=1, MajorId=1）",
        ["年份", "分数"],
        [("2023", "360"), ("2024", "370"), ("2025", "358")],
        [("ScoreLines", scoreline_fields)],
        screenshots=[],
    )

    # 图1-9
    add_section(
        doc,
        "图1-9",
        "智能志愿推荐界面图",
        "/recommendations",
        [
            "输入已录入成绩的用户 Id，可选年份（留空取该用户最近一条成绩年份）。",
            "可选志愿档位：全部 / 冲刺 / 匹配 / 保底。",
            "点击「生成推荐」：读取 ExamScores 总分，对比同年 ScoreLines，按分差划分档位并列出建议。",
        ],
        "输入模拟数据",
        ["输入项", "模拟值"],
        [
            ("用户 Id", "1（张三，2025 年总分 385）"),
            ("年份", "留空或 2025"),
            ("志愿档位", "全部档位"),
        ],
        None,
        screenshots=["fig1-9-recommendations.png"],
    )
    doc.add_heading("推荐结果示例", level=2)
    add_table(
        doc,
        ["院校", "专业", "分数线", "分差", "档位"],
        [
            ("北京大学", "计算机科学与技术", "358", "+27", "保底"),
            ("浙江大学", "计算机科学与技术", "355", "+30", "保底"),
            ("华中科技大学", "计算机科学与技术", "338", "+47", "保底"),
        ],
    )
    doc.add_heading("数据库表", level=2)
    doc.add_paragraph(
        "ExamScores（UserId, Year, TotalScore）、ScoreLines（Year, Score, SchoolId, MajorId）、"
        "Schools、Majors（展示名称与属性）"
    )

    # 图1-10
    add_section(
        doc,
        "图1-10",
        "招生简章查询界面图",
        "/recruitinfos",
        [
            "按院校 Id、专业 Id、年份筛选。",
            "点击「开始查询」，列表展示 RecruitInfos 及关联院校、专业名。",
            "点击条目，右侧展示招生计划、考试科目、其他要求、原文链接。",
        ],
        "模拟数据",
        ["列表项", "计划人数"],
        [
            ("2025 · 北京大学 · 计算机科学与技术", "35"),
            ("2025 · 华中科技大学 · 计算机科学与技术", "60"),
            ("2025 · 南京师范大学 · 学科教学（语文）", "30"),
            ("2024 · 清华大学 · 电子信息", "48"),
        ],
        [("RecruitInfos", recruit_fields)],
        screenshots=["fig1-10-recruitinfos.png"],
    )

    # 图1-11
    add_section(
        doc,
        "图1-11",
        "AI 助手对话界面图",
        "/ai-chat",
        [
            "用户输入自然语言问题（或点击示例问题填入输入框）。",
            "前端将对话上下文发送至后端 OpenAI 兼容接口（需配置 AiApiKey）。",
            "后端可通过检索服务引用院校、专业、分数线等数据生成回答。",
            "对话以气泡形式展示；当前实现不持久化聊天记录。",
        ],
        "模拟数据",
        ["类型", "内容"],
        [
            ("示例问题1", "计算机专硕国家线近年趋势？"),
            ("示例问题2", "如何根据分数选冲刺和保底院校？"),
            ("用户输入", "考研报名流程里要注意什么？"),
            ("AI 回复", "由接口实时生成，非固定模拟文本"),
        ],
        [
            (
                "引用数据表（无专用聊天表）",
                [
                    ("Schools", "院校 grounding"),
                    ("Majors", "专业 grounding"),
                    ("ScoreLines", "分数线 grounding"),
                ],
            ),
        ],
        screenshots=["fig1-11-aichat.png"],
    )

    # 图1-12
    add_section(
        doc,
        "图1-12",
        "管理后台界面图",
        "/admin（需 Admin 或 Root 角色）",
        [
            "新建用户：填写用户名、邮箱、初始密码，创建 Role=User 的普通用户。",
            "批量导入成绩：下载 Excel 模板 → 填写后上传 → 批量写入 ExamScores。",
            "用户列表：刷新、编辑、删除普通用户。",
            "仅 Root：创建管理员（Role=Admin）、管理员列表、对管理员「撤权」降为 User。",
        ],
        "新建用户模拟数据",
        ["用户名", "邮箱", "初始密码", "角色"],
        [("student_001", "student001@example.com", "Welcome2024!", "User")],
        None,
        screenshots=[
            "fig1-12-admin-create.png",
            "fig1-12-admin-users.png",
            "fig1-12-admin-root.png",
        ],
        screenshot_note="管理后台含三种典型视图：新建用户与 Excel 导入、普通用户列表、Root 管理员管理。",
    )
    doc.add_heading("Excel 导入一行示例", level=2)
    add_table(
        doc,
        ["用户Id", "年份", "总分", "政治", "英语", "数学", "专业课", "院校Id", "专业Id"],
        [("1", "2025", "385", "75", "78", "132", "100", "1", "1")],
    )
    doc.add_heading("用户列表示例", level=2)
    add_table(
        doc,
        ["Id", "用户名", "邮箱", "角色", "创建时间(UTC)"],
        [
            ("1", "张三", "zhangsan@test.com", "User", "2026-03-23 15:45"),
            ("2", "李四", "lisi@test.com", "User", "2026-03-23 15:45"),
            ("3", "王五", "wangwu@test.com", "Admin", "2026-03-23 15:45"),
        ],
    )
    doc.add_heading("Root 创建管理员示例", level=2)
    add_table(
        doc,
        ["用户名", "邮箱", "初始密码"],
        [("dept_manager", "manager@kyinfo.com", "SecurePass789")],
    )
    doc.add_heading("数据库表", level=2)
    add_table(doc, ["字段名", "说明"], user_fields)
    doc.add_paragraph("表名：Users")
    add_table(doc, ["字段名", "说明"], exam_fields)
    doc.add_paragraph("表名：ExamScores")
    add_table(doc, ["字段名", "说明"], audit_fields)
    doc.add_paragraph("表名：AuditLogs（管理操作审计）")

    # 图1-13
    add_section(
        doc,
        "图1-13",
        "我的成绩维护界面图",
        "/examscores",
        [
            "支持按任意用户 Id 查询成绩（便于管理员或自查）。",
            "可筛选年份、院校、专业并新增记录。",
            "功能与「我的档案」中成绩区类似，但查询范围不限于当前登录用户。",
        ],
        "模拟数据",
        ["用户 Id", "年份", "总分", "院校 Id", "专业 Id"],
        [("1", "2025", "385", "1", "1")],
        [("ExamScores", exam_fields)],
        screenshots=["fig1-4-account.png"],
        screenshot_note="独立路由为 /examscores；界面与「我的档案」中成绩维护区域相近，故沿用我的档案页截图示意。",
    )

    doc.add_page_break()
    doc.add_heading("界面与路由对照总表", level=1)
    add_table(
        doc,
        ["图号", "界面名称", "路由", "主要数据表"],
        [
            ("图1-1", "系统登录界面", "/login", "Users"),
            ("图1-2", "系统注册界面", "/register", "Users"),
            ("图1-3", "系统首页界面", "/", "—"),
            ("图1-4", "我的档案界面", "/account", "Users, ExamScores"),
            ("图1-5", "院校查询界面", "/schools", "Schools"),
            ("图1-6", "专业查询界面", "/majors", "Majors, Schools"),
            ("图1-7", "分数线查询界面", "/scorelines", "ScoreLines"),
            ("图1-8", "分数线趋势界面", "/scoreline-trends", "ScoreLines"),
            ("图1-9", "智能志愿推荐界面", "/recommendations", "ExamScores, ScoreLines, Schools, Majors"),
            ("图1-10", "招生简章查询界面", "/recruitinfos", "RecruitInfos, Schools, Majors"),
            ("图1-11", "AI 助手对话界面", "/ai-chat", "引用 Schools/Majors/ScoreLines"),
            ("图1-12", "管理后台界面", "/admin", "Users, ExamScores, AuditLogs"),
            ("图1-13", "我的成绩维护界面", "/examscores", "ExamScores"),
        ],
    )
    doc.add_paragraph(
        "说明：模拟数据与项目内 seed_kyinfo_testdata.sql 及系统截图展示保持一致，"
        "便于本地 https://localhost:7011 联调复现。"
    )

    out_path = r"D:\KyInfo.docx"
    doc.save(out_path)
    print(f"Saved: {out_path}")


if __name__ == "__main__":
    main()
